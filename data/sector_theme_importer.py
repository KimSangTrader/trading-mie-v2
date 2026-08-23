"""
SectorThemeImporter - 사용자가 1차 분류한 Sector/Theme 매핑 엑셀을 DB로 임포트 (Phase 5-12)

================================================================================
【변경 이력】
================================================================================
【2026-08-23】최초 생성

- 배경: SectorAnalyzer(가중치 0.18)/ThemeAnalyzer(가중치 0.09)가 여태 완전히 mock
  데이터로만 동작했음(Phase 5 완성도 점검 때 확인). KRX 공식 업종지수 API로
  연동하려던 조사(전날 세션, data/sector_code_master.py)는 SectorAnalyzer의 8개
  카테고리가 KRX 공식 업종분류와 1:1로 안 맞고, 매핑을 사람이 판단해야 하는
  문제에 부딪혔다. 사용자가 그 판단을 직접 해서(KRX_전체종목_Sector_Theme_매핑.xlsx,
  종목 2,803개 - 유가/코스닥/코넥스 전체) 1차 분류를 엑셀로 만들어 제공 - 이 모듈은
  그 엑셀을 DB(stock_sector_mapping/stock_theme_mapping, db/models.py 참고)에
  저장하는 역할만 담당한다. SectorAnalyzer/ThemeAnalyzer가 이 DB를 실제로 어떻게
  써서 점수를 낼지(예: 섹터별 종목 바스켓의 실시간 등락률 평균 등)는 별도 논의
  필요 - 이 모듈 범위 밖.

- 엑셀 구조(전체종목_매핑 시트, 2,803행):
  * 종목코드/종목명/시장구분(유가=KOSPI/코스닥=KOSDAQ/코넥스=KONEX)/
    KRX_원본업종/주요제품 - 종목 기본 정보
  * Analysis_Sector - SectorAnalyzer용 카테고리(12종: IT_Semiconductor,
    Healthcare_Pharma, Telecom_Media, Industrials, Finance, Chemicals_Energy,
    Consumer, Construction_Real_Estate, Materials, Utilities, Transportation,
    Other) - 기존 SectorAnalyzer가 하드코딩하고 있던 8개(Secondary_Battery
    포함)와 다르다. Secondary_Battery는 여기 없고 Theme 쪽에 있음(전날 세션에서
    예상했던 대로 - 2차전지는 KRX 전통 업종이 아니라 여러 업종에 걸친 테마임이
    이걸로 확인됨). SectorAnalyzer의 self.sectors 하드코딩 딕셔너리는 이 실제
    카테고리에 맞춰 별도로 업데이트해야 한다(이 모듈 범위 밖, 후속 작업).
  * Primary_Theme(0~1개)/Secondary_Themes(0~N개, 콤마로 구분) - ThemeAnalyzer용,
    18종(Healthcare/Bio/Semiconductor/Automotive/Content/Food/
    Internet_Platform/Secondary_Battery/K_Beauty/AI/Robotics/Shipbuilding/
    Renewable_Energy/Gaming/REITs/Defense/EV/Nuclear) - 기존 ThemeAnalyzer의
    6개 추상적 "시장 심리" 카테고리(geopolitical_risk 등)와는 성격이 완전히
    다르다(종목별 실제 소속 테마 태그) - ThemeAnalyzer도 이 실제 구조에 맞춰
    재설계가 필요할 가능성이 높다(이 모듈 범위 밖, 후속 논의 필요).
  * 검토필요여부(Y/N) - 자동판정 신뢰도가 낮아 사람 확인이 필요하다고 표시된
    행(1,056건, 37.7%) - DB에 needs_review로 그대로 보존해서, 나중에 분석에서
    이 값을 제외할지 포함할지 소비하는 쪽(SectorAnalyzer 등)이 정책을 정할 수
    있게 한다(이 모듈은 필터링하지 않고 있는 그대로 저장).
  * '검토필요' 시트는 전체종목_매핑에서 검토필요여부='Y'인 행만 뽑아둔 부분집합
    이라(헤더/행 수 확인함) 별도로 임포트하지 않는다 - 소스 오브 트루스는
    전체종목_매핑 하나뿐.
  * 완전히 동일한 내용으로 중복된 행이 43건 있었다(예: 000480 CR홀딩스 2회,
    Sector/Theme 값 포함 전부 동일 - 서로 다른 값으로 충돌하는 중복은 0건임을
    확인함). 파싱 시 종목코드 기준으로 중복 제거(첫 값 유지)한다.

- DB 저장 방식: stock_valuation과 동일한 "배치(timestamp)" 패턴을 그대로 따른다 -
  임포트할 때마다 새 timestamp로 통째로 insert하고(UPDATE 아님), 분석은 항상
  MAX(timestamp) 배치만 읽는다(get_latest_sector_mapping/get_latest_theme_mapping
  참고). 사용자가 나중에 분류를 다시 정리해서 새 엑셀을 주면 이 스크립트를 다시
  실행하기만 하면 되고, 과거 배치는 지우지 않아 이력이 남는다.

- 이 세션은 실제 RDS에 접속해 임포트를 검증할 수 없었다(클라우드 샌드박스 네트워크
  제약, 지금까지와 동일한 한계) - 다만 사용자가 첨부한 실제 엑셀 파일 자체는 이
  세션에서 읽을 수 있어서, parse_mapping_excel()의 파싱/정규화/중복제거 로직은
  실제 파일로 직접 검증했다(첨부 파일 기준 2,803행 → 중복 43건 제거 → 2,760종목,
  시장 정규화 3종 전부 매핑 확인, Sector 12종/Theme 18종 전부 확인). DB insert
  부분(import_to_db)만 SQLite로 별도 검증했고, 실제 PostgreSQL RDS 연결은
  사용자 컴퓨터/서버에서 라이브 확인 필요.
================================================================================
"""

import logging
import os
import sys
from collections import OrderedDict
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

# 【2026-08-23 추가】`python data\sector_theme_importer.py`처럼 파일 경로로 직접
# 실행하면 파이썬이 이 파일이 있는 data/ 폴더를 sys.path[0]으로 잡아버려서,
# 프로젝트 루트에 있는 config/db 패키지를 못 찾는다("ModuleNotFoundError: No
# module named 'config'" - 라이브 실행 중 실제로 재현됨). 이 파일은 data/ 바로
# 아래에 있으므로 부모 디렉터리(프로젝트 루트)를 sys.path 맨 앞에 추가해서,
# `python data\sector_theme_importer.py`로 직접 실행하든 프로젝트 루트에서
# `python -m data.sector_theme_importer`로 실행하든 항상 동작하게 한다.
_PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

logger = logging.getLogger(__name__)

DEFAULT_SHEET_NAME = "전체종목_매핑"

# 엑셀 원본 시장구분 라벨 -> 이 프로젝트 전반에서 쓰는 정규화된 시장 코드
# (data/stock_master.py 등 기존 코드는 'KOSPI'/'KOSDAQ' 영문 상수를 씀)
_MARKET_MAP = {
    "유가": "KOSPI",
    "코스닥": "KOSDAQ",
    "코넥스": "KONEX",
}

# 이 컬럼들이 헤더에 없으면 엑셀 형식이 바뀐 것으로 보고 즉시 예외를 던진다
# (조용히 잘못된 컬럼 위치로 파싱하지 않기 위한 방어 장치 - stock_master.py의
# _MIN_EXPECTED_RECORDS 자기검증과 같은 원칙)
_REQUIRED_COLUMNS = [
    "종목코드", "종목명", "시장구분", "KRX_원본업종", "주요제품",
    "Analysis_Sector", "Primary_Theme", "Secondary_Themes",
    "검토필요여부", "Sector_판정방식", "Theme_판정방식",
]


def _normalize_market(raw_market: Optional[str], ticker: str) -> str:
    market = _MARKET_MAP.get((raw_market or "").strip())
    if market is None:
        raise ValueError(
            f"알 수 없는 시장구분 '{raw_market}' (종목코드={ticker}) - "
            f"엑셀 형식이 바뀌었을 수 있습니다. _MARKET_MAP을 확인/보강하세요."
        )
    return market


def _split_secondary_themes(raw: Optional[str]) -> List[str]:
    if not raw:
        return []
    return [t.strip() for t in str(raw).split(",") if t.strip()]


def parse_mapping_excel(file_path: str, sheet_name: str = DEFAULT_SHEET_NAME) -> Dict[str, Any]:
    """엑셀을 읽어 DB insert용 dict 리스트로 변환한다 (DB 세션 없이 순수 파싱만 -
    독립적으로 테스트 가능).

    Returns:
        {
            "sector_rows": [{"ticker":.., "name":.., "market":.., "market_raw":..,
                              "krx_industry":.., "main_products":.., "sector":..,
                              "needs_review":.., "sector_method":.., "theme_method":..}, ...],
            "theme_rows": [{"ticker":.., "theme":.., "is_primary":..}, ...],
            "stats": {"total_rows": int, "duplicate_rows_dropped": int,
                      "unique_tickers": int, "needs_review_count": int,
                      "sector_counts": {...}, "theme_counts": {...}},
        }
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"시트 '{sheet_name}'을 찾을 수 없습니다. 실제 시트 목록: {wb.sheetnames}"
        )
    ws = wb[sheet_name]

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = list(header_row)
    missing = [c for c in _REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise ValueError(
            f"엑셀에 필수 컬럼이 없습니다: {missing} (실제 헤더: {headers}) - "
            f"엑셀 형식이 바뀌었을 수 있습니다."
        )
    idx = {name: i for i, name in enumerate(headers)}

    # 종목코드 기준 중복 제거 (완전 동일한 행만 존재함을 확인했으므로 첫 값 유지 -
    # 값이 서로 다른 충돌 중복이 생기면 발견 즉시 알 수 있도록 아래에서 검증한다)
    seen: "OrderedDict[str, Tuple]" = OrderedDict()
    conflicts: List[str] = []
    total_rows = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["종목코드"]] is None:
            continue  # 완전 빈 줄
        total_rows += 1
        ticker = str(row[idx["종목코드"]]).strip()

        key_fields = (
            row[idx["Analysis_Sector"]], row[idx["Primary_Theme"]],
            row[idx["Secondary_Themes"]], row[idx["시장구분"]],
        )
        if ticker in seen:
            if seen[ticker][1] != key_fields:
                conflicts.append(ticker)
            continue  # 첫 값 유지 (완전 동일 중복이면 그대로 스킵)
        seen[ticker] = (row, key_fields)

    if conflicts:
        raise ValueError(
            f"종목코드가 같은데 Sector/Theme 값이 서로 다른 중복 행이 있습니다: "
            f"{conflicts[:10]}{'...' if len(conflicts) > 10 else ''} - "
            f"어느 값이 맞는지 사람이 확인해야 하므로 자동으로 고르지 않습니다."
        )

    duplicate_rows_dropped = total_rows - len(seen)

    sector_rows: List[Dict[str, Any]] = []
    theme_rows: List[Dict[str, Any]] = []
    sector_counts: Dict[str, int] = {}
    theme_counts: Dict[str, int] = {}
    needs_review_count = 0

    for ticker, (row, _key) in seen.items():
        market_raw = row[idx["시장구분"]]
        market = _normalize_market(market_raw, ticker)
        needs_review = str(row[idx["검토필요여부"]]).strip().upper() == "Y"
        sector = row[idx["Analysis_Sector"]]

        if needs_review:
            needs_review_count += 1
        if sector:
            sector_counts[sector] = sector_counts.get(sector, 0) + 1

        sector_rows.append({
            "ticker": ticker,
            "name": row[idx["종목명"]],
            "market": market,
            "market_raw": market_raw,
            "krx_industry": row[idx["KRX_원본업종"]],
            "main_products": row[idx["주요제품"]],
            "sector": sector,
            "needs_review": needs_review,
            "sector_method": row[idx["Sector_판정방식"]],
            "theme_method": row[idx["Theme_판정방식"]],
        })

        primary_theme = row[idx["Primary_Theme"]]
        if primary_theme:
            theme_rows.append({"ticker": ticker, "theme": primary_theme, "is_primary": True})
            theme_counts[primary_theme] = theme_counts.get(primary_theme, 0) + 1

        for theme in _split_secondary_themes(row[idx["Secondary_Themes"]]):
            theme_rows.append({"ticker": ticker, "theme": theme, "is_primary": False})
            theme_counts[theme] = theme_counts.get(theme, 0) + 1

    return {
        "sector_rows": sector_rows,
        "theme_rows": theme_rows,
        "stats": {
            "total_rows": total_rows,
            "duplicate_rows_dropped": duplicate_rows_dropped,
            "unique_tickers": len(sector_rows),
            "needs_review_count": needs_review_count,
            "sector_counts": sector_counts,
            "theme_counts": theme_counts,
        },
    }


def import_to_db(file_path: str, session, sheet_name: str = DEFAULT_SHEET_NAME) -> Dict[str, Any]:
    """parse_mapping_excel() 결과를 stock_sector_mapping/stock_theme_mapping에
    새 배치(timestamp)로 insert하고 커밋한다. 기존 배치는 건드리지 않는다(이력 보존)."""
    from db.models import StockSectorMapping, StockThemeMapping

    parsed = parse_mapping_excel(file_path, sheet_name)
    batch_ts = datetime.now(timezone.utc)

    for row in parsed["sector_rows"]:
        session.add(StockSectorMapping(timestamp=batch_ts, **row))
    for row in parsed["theme_rows"]:
        session.add(StockThemeMapping(timestamp=batch_ts, **row))

    session.commit()

    return {"batch_timestamp": batch_ts, **parsed["stats"]}


def get_latest_sector_mapping(session) -> List[Dict[str, Any]]:
    """최신 배치의 종목별 Sector 매핑을 dict 리스트로 반환 (stock_valuation의
    get_latest_stock_valuations()와 동일한 "최신 timestamp만" 패턴)."""
    from sqlalchemy import func
    from db.models import StockSectorMapping

    latest_ts = session.query(func.max(StockSectorMapping.timestamp)).scalar()
    if latest_ts is None:
        return []

    rows = (
        session.query(StockSectorMapping)
        .filter(StockSectorMapping.timestamp == latest_ts)
        .all()
    )
    return [
        {
            "ticker": r.ticker, "name": r.name, "market": r.market,
            "sector": r.sector, "needs_review": r.needs_review,
        }
        for r in rows
    ]


def get_latest_theme_mapping(session) -> List[Dict[str, Any]]:
    """최신 배치의 종목별 Theme 매핑을 dict 리스트로 반환."""
    from sqlalchemy import func
    from db.models import StockThemeMapping

    latest_ts = session.query(func.max(StockThemeMapping.timestamp)).scalar()
    if latest_ts is None:
        return []

    rows = (
        session.query(StockThemeMapping)
        .filter(StockThemeMapping.timestamp == latest_ts)
        .all()
    )
    return [
        {"ticker": r.ticker, "theme": r.theme, "is_primary": r.is_primary}
        for r in rows
    ]


if __name__ == "__main__":
    import sys

    logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")

    if len(sys.argv) < 2:
        print("사용법: python data/sector_theme_importer.py <엑셀파일경로>")
        sys.exit(1)

    excel_path = sys.argv[1]

    from config.database import SessionLocal, init_db

    print(f"\n【1단계】테이블 생성 확인 (없으면 생성)...")
    init_db()

    print(f"\n【2단계】엑셀 파싱 중... ({excel_path})")
    session = SessionLocal()
    try:
        result = import_to_db(excel_path, session)
    finally:
        session.close()

    stats = {k: v for k, v in result.items() if k != "sector_counts" and k != "theme_counts"}
    print(f"\n✅ 임포트 완료 - {stats}")
    print(f"\n【Sector 분포】")
    for sector, count in sorted(result["sector_counts"].items(), key=lambda x: -x[1]):
        print(f"  {sector:28} {count:5}종목")
    print(f"\n【Theme 분포】")
    for theme, count in sorted(result["theme_counts"].items(), key=lambda x: -x[1]):
        print(f"  {theme:28} {count:5}종목(태그 수, 중복 포함)")
